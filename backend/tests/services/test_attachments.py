"""附件功能单测：文本抽取（txt/docx/xlsx/pdf）、上下文拼装、暂存生命周期。"""

import pytest

from app.services.attachment_service import (
    AttachmentError,
    attachment_context,
    content_with_attachments,
    extract_text,
    extra_with_attachments,
    load_staged,
    save_staging,
)


def test_extract_text_plain():
    assert extract_text("a.txt", "营收 100 万".encode()) == "营收 100 万"


def test_extract_docx(tmp_path):
    import io

    from docx import Document

    doc = Document()
    doc.add_paragraph("第一季度营收 300 万")
    doc.add_paragraph("第二季度营收 450 万")
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("季报.docx", buf.getvalue())
    assert "第一季度营收 300 万" in text and "第二季度营收 450 万" in text


def test_extract_xlsx_rows_and_sheet_names():
    import io

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "销售"
    ws.append(["月份", "销售额"])
    ws.append(["1月", 100])
    ws.append(["2月", 200])
    text = extract_text("销售.xlsx", _xlsx_bytes(wb))
    assert "# 工作表: 销售" in text
    assert "月份 | 销售额" in text and "1月 | 100" in text


def _xlsx_bytes(wb) -> bytes:
    import io

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_extract_pdf():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    # 空白页无文本：验证不崩溃即可；正文抽取由 e2e 覆盖
    assert isinstance(extract_text("blank.pdf", _pdf_bytes(writer)), str)


def _pdf_bytes(writer) -> bytes:
    import io

    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def test_extract_rejects_unsupported():
    with pytest.raises(AttachmentError, match="不支持的附件类型"):
        extract_text("app.bin", b"\x00\x01")


def test_extract_truncates_per_file():
    text = extract_text("big.txt", ("字" * 30_000).encode())
    assert len(text) == 15_000


def test_content_with_attachments_keeps_clean_content():
    atts = [{"name": "a.txt", "kind": "text", "size": 5, "text": "关键数据"}]
    out = content_with_attachments("总结附件", atts)
    assert out.startswith("总结附件")
    assert "[附件 a.txt" in out and "关键数据" in out
    assert content_with_attachments("总结附件", None) == "总结附件"
    assert content_with_attachments("总结附件", [{"name": "空.txt", "kind": "text", "size": 0, "text": ""}]) == "总结附件"


def test_attachment_context_total_budget():
    atts = [
        {"name": f"{i}.txt", "kind": "text", "size": 1, "text": "x" * 15_000}
        for i in range(5)
    ]
    ctx = attachment_context(atts)
    assert ctx.count("[附件") == 5  # 第 4 份起只剩省略占位也算一条
    assert len(ctx) <= 50_000 + 5 * 60  # 总量封顶 + 每份头部说明


def test_extra_merge_preserves_route():
    atts = [{"name": "a.txt", "kind": "text", "size": 5, "text": "t"}]
    merged = extra_with_attachments(atts, {"route": {"status": "pending"}})
    assert merged["route"] == {"status": "pending"}
    assert merged["attachments"] == atts
    assert extra_with_attachments(None, {"route": {}}) == {"route": {}}
    assert extra_with_attachments(None) is None


def test_staging_roundtrip_one_shot(tmp_path, monkeypatch):
    from app.config import get_settings

    monkeypatch.setenv("UPLOAD_ROOT", str(tmp_path / "up"))
    get_settings.cache_clear()
    meta = save_staging("u1", "数据.csv", "a,b\n1,2".encode())
    assert meta["kind"] == "csv" and "a,b" in meta["text"]

    loaded = load_staged("u1", [meta["id"]])
    assert loaded[0]["name"] == "数据.csv"
    # 取回即删：重复引用报错
    with pytest.raises(AttachmentError, match="重新上传"):
        load_staged("u1", [meta["id"]])
    get_settings.cache_clear()


def test_staging_enforces_size_limit(tmp_path, monkeypatch):
    from app.config import get_settings

    monkeypatch.setenv("UPLOAD_ROOT", str(tmp_path / "up"))
    get_settings.cache_clear()
    with pytest.raises(AttachmentError, match="20MB"):
        save_staging("u1", "huge.txt", b"x" * (20 * 1024 * 1024 + 1))
    get_settings.cache_clear()
