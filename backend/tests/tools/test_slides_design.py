"""幻灯片设计系统单测：主题/版式归一、三明治补齐、数字卡片、pptx 有效性。"""

import io
import zipfile

from app.tools.slides import (
    DEFAULT_THEME,
    THEMES,
    body_mode,
    build_slides_html,
    build_slides_pptx,
    normalize_slides,
    resolve_theme,
    resolve_theme_name,
    split_unit,
    stat_parts,
)


def test_stat_parts_matches_and_strips_connectors():
    assert stat_parts("37% 复购率提升") == ("37%", "复购率提升")
    assert stat_parts("35% 的核心客群为 25-35 岁") == ("35%", "核心客群为 25-35 岁")
    assert stat_parts("120 家门店") == ("120家", "门店")  # 量词并入大数字
    assert stat_parts("28元 客单价") == ("28元", "客单价")
    assert stat_parts("8 亿 三年目标营收规模") == ("8亿", "三年目标营收规模")  # 单位前不留空格
    assert stat_parts("核心用户 8 万") is None  # 非数字开头不匹配
    assert stat_parts("纯文字要点") is None


def test_split_unit_separates_digits_from_suffix():
    assert split_unit("1,500亿") == ("1,500", "亿")
    assert split_unit("37%") == ("37", "%")
    assert split_unit("120家") == ("120", "家")
    assert split_unit("28元") == ("28", "元")
    assert split_unit("22") == ("22", "")  # 无单位不拆


def test_resolve_theme_falls_back_to_default():
    assert resolve_theme_name("teal") == "teal"
    assert resolve_theme_name("not-a-theme") == DEFAULT_THEME
    assert resolve_theme_name(None) == DEFAULT_THEME
    assert resolve_theme("coral") is THEMES["coral"]


def test_normalize_adds_cover_and_end_sandwich():
    norm = normalize_slides("季度汇报", [{"title": "进展", "bullets": ["A"]}])
    assert [s["layout"] for s in norm] == ["cover", "content", "end"]
    assert norm[0]["title"] == "季度汇报"  # 封面用 deck 标题
    assert norm[-1]["subtitle"] == "季度汇报"


def test_normalize_respects_explicit_cover_and_numbers_sections():
    norm = normalize_slides(
        "T",
        [
            {"layout": "cover", "title": "封面"},
            {"layout": "section", "title": "一"},
            {"layout": "content", "title": "页", "bullets": ["x"]},
            {"layout": "section", "title": "二"},
            {"layout": "end", "title": "谢"},
        ],
    )
    assert [s["layout"] for s in norm] == ["cover", "section", "content", "section", "end"]
    assert [s["no"] for s in norm if s["layout"] == "section"] == ["01", "02"]


def test_normalize_coerces_dirty_layout_and_keeps_input_intact():
    raw = [{"title": "p", "bullets": ["a"], "layout": "weird"}]
    norm = normalize_slides("T", raw)
    assert norm[0]["layout"] == "cover"  # 自动补的封面在前
    assert norm[1]["layout"] == "content"  # 脏 layout 归一
    assert raw[0]["layout"] == "weird"  # 入参不被改写


def test_html_renders_theme_tokens_and_stat_cards():
    html = build_slides_html(
        "发布会",
        [{"title": "增长", "bullets": ["37% 复购率提升", "核心用户 8 万", "35% 的核心客群为 25-35 岁"]}],
        "coral",
    )
    assert THEMES["coral"]["dark"].lower() in html.lower()
    assert 'class="num"' in html  # 数字卡片分支
    # 徽章贯穿所有卡片——数字卡片不得顶掉序号(视觉 QA 修复)
    assert html.count('class="badge"') == 3
    assert ">核心客群为" in html  # 说明剥掉“的”连接字
    assert "的核心客群" not in html


def test_html_many_bullets_falls_back_to_list_mode():
    bullets = [f"要点{i}" for i in range(9)]
    html = build_slides_html("T", [{"title": "p", "bullets": bullets}], "teal")
    assert 'class="list"' in html and 'class="cards"' not in html


def test_body_mode_dispatch_by_stat_ratio_and_count():
    assert body_mode(["19% 增速", "1,500亿 规模", "22杯 人均"]) == "kpis"  # 全数据页
    assert body_mode(["a", "b", "c", "d"]) == "grid2"  # ≥4 条纯文字 → 两列
    assert body_mode(["37% 提升", "纯文字要点"]) == "cards"  # 混合 → 纵向卡片
    assert body_mode([f"b{i}" for i in range(8)]) == "list"  # 防溢出
    assert body_mode([]) == "cards"


def test_html_kpi_grid_for_all_stat_page():
    html = build_slides_html(
        "T",
        [{"title": "市场", "bullets": ["19% 年复合增速", "1,500亿 市场规模",
                                       "22杯 人均年饮用", "35% 年轻客群占比"]}],
        "midnight",
    )
    assert 'class="kpis"' in html
    assert html.count('class="kpi"') == 4
    assert 'class="badge"' not in html  # KPI 网格不叠序号:数字即锚点
    assert "repeat(2,1fr)" in html  # 4 条 → 2×2
    assert '>19<span class="u">%</span>' in html  # 单位小号化,数字是主角


def test_html_cover_renders_stat_hooks():
    html = build_slides_html(
        "咖啡计划",
        [{"layout": "cover", "title": "咖啡计划", "bullets": ["120家 三年门店", "8亿 目标营收"]}],
        "teal",
    )
    assert html.count('class="hook"') == 2  # 封面数字钩子
    assert '>120<span class="u">家</span></div>' in html  # 量词小号化


def test_html_agenda_layout_numbers_items():
    html = build_slides_html("T", [{"layout": "agenda", "title": "目录", "bullets": ["市场", "财务"]}], "teal")
    assert 'class="item"' in html and ">01<" in html and ">市场<" in html


def test_pptx_bytes_is_valid_package_with_sandwich():
    data = build_slides_pptx(
        "规划",
        [
            {"layout": "cover", "title": "规划", "subtitle": "2026"},
            {"layout": "section", "title": "展望"},
            {"title": "目标", "bullets": ["三条线并进", "50% 增速目标"]},
        ],
        "forest",
    )
    assert zipfile.is_zipfile(io.BytesIO(data))  # OOXML 本质是 zip

    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 4  # 显式封面 + section + content + 自动结尾
    assert prs.slide_width == 9144000 and prs.slide_height == 5143500  # 16:9


def test_pptx_many_bullets_uses_compact_list():
    data = build_slides_pptx("T", [{"title": "p", "bullets": [f"b{i}" for i in range(9)]}], None)
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    texts = [sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame]
    assert any("· b0" in t for t in texts)


def test_pptx_kpi_and_agenda_pages_valid():
    data = build_slides_pptx(
        "T",
        [
            {"layout": "agenda", "title": "目录", "bullets": ["市场", "财务", "团队"]},
            {"title": "市场", "subtitle": "市场机会", "bullets": ["19% 年复合增速", "1,500亿 市场规模", "22杯 人均"]},
        ],
        "coral",
    )
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 4  # 自动封面 + agenda + content + 自动结尾
    kpi_texts = [sh.text_frame.text for sh in prs.slides[2].shapes if sh.has_text_frame]
    assert any("19%" in t for t in kpi_texts) and any("市场机会" in t for t in kpi_texts)
