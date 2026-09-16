"""产物工具单测：文件落盘、格式可读、降级路径、版本自增、上下文缺失。"""

import io
import json

import pytest

from app.storage.local_disk import LocalDiskStorage
from app.tools.artifacts import register_artifact_tools
from app.tools.base import ToolError, ToolRegistry
from app.tools.context import ToolContext, set_tool_context

USER, TASK = "u1", "t1"


@pytest.fixture
def registry(storage) -> ToolRegistry:
    reg = ToolRegistry()
    register_artifact_tools(reg)
    set_tool_context(ToolContext(user_id=USER, task_id=TASK, storage=storage))
    return reg


def _result(raw: str) -> dict:
    return json.loads(raw)


def _keys(result: dict) -> list[str]:
    return [s["key"] for s in result["saved"]]


async def test_save_document_writes_md_and_docx(registry, storage):
    raw = await registry.execute(
        "save_document",
        {"title": "周报", "content_md": "## 进展\n- 完成工具层\n\n正文段落"},
    )
    result = _result(raw)
    md_key = _keys(result)[0]
    assert md_key == f"{USER}/{TASK}/v1/document.md"
    assert "完成工具层" in storage.read_text(md_key)

    from docx import Document

    doc = Document(io.BytesIO(storage.read(f"{USER}/{TASK}/v1/document.docx")))
    texts = [p.text for p in doc.paragraphs]
    assert "周报" in texts and "完成工具层" in texts


async def test_save_document_rejects_empty(registry):
    with pytest.raises(ToolError, match="文档内容为空"):
        await registry.execute("save_document", {"title": "空", "content_md": "  "})


async def test_save_poster_html_and_png_degrade(registry, storage, monkeypatch):
    def _no_renderer(*a, **kw):
        raise RuntimeError("chromium 未安装")

    monkeypatch.setattr("app.tools.artifacts._render_png", _no_renderer)
    result = _result(
        await registry.execute(
            "save_poster",
            {"title": "新品发布会<script>", "slogan": "9月见"},
        )
    )
    key = _keys(result)[0]
    html = storage.read_text(key)
    assert "新品发布会&lt;script&gt;" in html  # XSS：标题经转义
    assert "跳过" in result["note"]  # 无浏览器时降级，不失败


async def test_save_poster_png_success(registry, storage, monkeypatch):
    def _fake_png(html_key: str, filename: str) -> str:
        png_key = f"{html_key.rsplit('/', 1)[0]}/{filename}"
        storage.save(png_key, b"\x89PNG fake")
        return png_key

    monkeypatch.setattr("app.tools.artifacts._render_png", _fake_png)
    result = _result(await registry.execute("save_poster", {"title": "海报"}))
    assert "poster.png" in result["note"]
    assert storage.exists(f"{USER}/{TASK}/v1/poster.png")  # png 与 html 同版本目录


async def test_save_table_xlsx_and_html_preview(registry, storage):
    result = _result(
        await registry.execute(
            "save_table",
            {
                "title": "销量<b>表</b>",
                "headers": ["月份", "销量"],
                "rows": [["7月", 120], ["8月", 150]],
            },
        )
    )
    keys = _keys(result)
    xlsx_key = next(k for k in keys if k.endswith(".xlsx"))
    html_key = next(k for k in keys if k.endswith(".html"))

    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(storage.read(xlsx_key))).active
    assert [c.value for c in ws[1]] == ["月份", "销量"]
    assert ws.max_row == 3

    html = storage.read_text(html_key)
    assert "销量&lt;b&gt;表&lt;/b&gt;" in html  # XSS：单元格转义
    assert "<td>150</td>" in html


async def test_save_table_column_mismatch(registry):
    with pytest.raises(ToolError, match="列数"):
        await registry.execute(
            "save_table",
            {"title": "t", "headers": ["a", "b"], "rows": [["1"]]},
        )


async def test_save_slides_html_and_pptx(registry, storage):
    result = _result(
        await registry.execute(
            "save_slides",
            {
                "title": "述职",
                "slides": [
                    {"title": "背景", "bullets": ["业务扩张"]},
                    {"title": "成果", "bullets": ["收入 +30%", "NPS 提升"]},
                ],
            },
        )
    )
    html = storage.read_text(_keys(result)[0])
    # 自动三明治：2 内容页 + 自动封面/结尾 = 4 页
    assert html.count('class="slide') == 4
    assert 'class="slide cover"' in html and 'class="slide end"' in html
    assert "NPS 提升" in html
    assert "收入" in html

    from pptx import Presentation

    prs = Presentation(io.BytesIO(storage.read(f"{USER}/{TASK}/v1/slides.pptx")))
    assert len(prs.slides) == 4


async def test_save_slides_requires_pages(registry):
    with pytest.raises(ToolError, match="至少需要一页"):
        await registry.execute("save_slides", {"title": "空", "slides": []})


async def test_versions_increment_across_saves(registry, storage):
    await registry.execute("save_document", {"title": "a", "content_md": "一版"})
    await registry.execute("save_document", {"title": "a", "content_md": "二版"})
    assert storage.exists(f"{USER}/{TASK}/v1/document.md")
    assert storage.exists(f"{USER}/{TASK}/v1/document.docx")  # 一次保存同目录
    assert storage.exists(f"{USER}/{TASK}/v2/document.md")


async def test_context_missing_raises_tool_error():
    """未注入工具上下文（不在 runner 执行流内）时统一转 ToolError。"""
    from app.tools.context import clear_tool_context

    reg = ToolRegistry()
    register_artifact_tools(reg)
    clear_tool_context()  # 夹具在同线程上下文注入过，显式清除后再验证
    with pytest.raises(ToolError, match="上下文未设置"):
        await reg.execute("save_document", {"title": "x", "content_md": "y"})
