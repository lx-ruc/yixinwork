"""附件服务：上传暂存、文本抽取、拼装 LLM 上下文。

设计：上传时即抽取文本（失败趁早暴露给用户），暂存 JSON 于磁盘；
发送消息时合并进 message.extra.attachments（含文本），此后聊天历史
回放、路由确认、降级直答、工作模式任务共用这一份上下文。
"""

import json
import uuid
from pathlib import Path

from app.config import get_settings

MAX_FILES_PER_MESSAGE = 5
MAX_FILE_BYTES = 20 * 1024 * 1024  # 单文件 20MB
PER_FILE_CHARS = 15_000  # 单附件抽取文本截断
TOTAL_CHARS = 50_000  # 单条消息附件上下文总量
XLSX_MAX_ROWS = 200  # 表格类附件预览行数

PLAINTEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log", ".xml",
    ".yaml", ".yml", ".html", ".htm", ".py", ".js", ".ts", ".css", ".sql", ".ini",
}
SUPPORTED_SUFFIXES = PLAINTEXT_SUFFIXES | {".docx", ".xlsx", ".pdf"}

KIND_BY_SUFFIX = {
    ".docx": "word",
    ".xlsx": "excel",
    ".pdf": "pdf",
    ".csv": "csv",
    ".tsv": "csv",
}


class AttachmentError(Exception):
    """附件不可用（上传/引用阶段抛给用户，HTTP 4xx）。"""


def _kind(name: str) -> str:
    return KIND_BY_SUFFIX.get(Path(name).suffix.lower(), "text")


def extract_text(name: str, data: bytes) -> str:
    """按扩展名抽取文本；不支持的类型抛 AttachmentError。截断到 PER_FILE_CHARS。"""
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise AttachmentError(
            f"不支持的附件类型「{suffix or name}」，可用：文档(pdf/docx)、"
            "表格(xlsx/csv)、文本(txt/md/json/log/代码等)"
        )
    if suffix == ".docx":
        text = _extract_docx(data)
    elif suffix == ".xlsx":
        text = _extract_xlsx(data)
    elif suffix == ".pdf":
        text = _extract_pdf(data)
    else:
        text = data.decode("utf-8", errors="replace")
    return text.strip()[:PER_FILE_CHARS]


def _extract_docx(data: bytes) -> str:
    import io

    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"# 工作表: {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= XLSX_MAX_ROWS:
                parts.append(f"……（超出前 {XLSX_MAX_ROWS} 行省略）")
                break
            parts.append(" | ".join("" if v is None else str(v) for v in row))
    wb.close()
    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    import io

    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:  # noqa: BLE001 pypdf 对损坏文件抛各式异常
        raise AttachmentError(f"PDF 解析失败: {exc}") from exc


def attachment_context(attachments: list[dict]) -> str:
    """把附件列表拼成 LLM 可读的上下文块；预留总量截断。"""
    if not attachments:
        return ""
    blocks: list[str] = []
    used = 0
    for a in attachments:
        text = str(a.get("text", "")).strip()
        if not text:
            continue
        budget = max(0, min(PER_FILE_CHARS, TOTAL_CHARS - used))
        if budget == 0:
            blocks.append(f"[附件 {a.get('name')} 内容已省略：附件总量超限]")
            continue
        clipped = text[:budget]
        used += len(clipped)
        blocks.append(
            f"[附件 {a.get('name')}（{a.get('kind', 'text')}，{a.get('size', 0)} 字节）]\n"
            f"{clipped}\n[/附件]"
        )
    if not blocks:
        return ""
    return "\n\n" + "\n\n".join(blocks)


def content_with_attachments(content: str, attachments: list[dict] | None) -> str:
    """用户可见 content 保持干净；给 LLM 的版本在其后拼附件上下文。"""
    if not attachments:
        return content
    ctx = attachment_context(attachments)
    return f"{content}{ctx}" if ctx else content


def extra_with_attachments(
    attachments: list[dict] | None, base: dict | None = None
) -> dict | None:
    """附件元数据并入消息 extra；无附件时保持原样（不产生空键）。"""
    if not attachments:
        return base
    return {**(base or {}), "attachments": attachments}


def _staging_dir(user_id: str) -> Path:
    root = Path(get_settings().upload_root) / user_id
    root.mkdir(parents=True, exist_ok=True)
    return root


def save_staging(user_id: str, name: str, data: bytes) -> dict:
    """上传暂存：抽取文本后落 JSON，返回附件元数据（发送时引用 id 合并）。"""
    if len(data) > MAX_FILE_BYTES:
        raise AttachmentError(f"附件「{name}」超过 20MB 上限")
    if not data:
        raise AttachmentError(f"附件「{name}」为空文件")
    meta = {
        "id": uuid.uuid4().hex,
        "name": name,
        "size": len(data),
        "kind": _kind(name),
        "text": extract_text(name, data),
    }
    path = _staging_dir(user_id) / f"{meta['id']}.json"
    path.write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    return meta


def load_staged(user_id: str, ids: list[str]) -> list[dict]:
    """发送时按 id 取回暂存附件（校验归属）；取回即删，防止重复引用。"""
    metas: list[dict] = []
    staging = _staging_dir(user_id)
    for attachment_id in ids:
        path = staging / f"{attachment_id}.json"
        if not path.is_file():
            raise AttachmentError(f"附件不存在或已使用，请重新上传（id={attachment_id}）")
        metas.append(json.loads(path.read_text(encoding="utf-8")))
        path.unlink(missing_ok=True)
    return metas


def sweep_staging(user_id: str, max_age_seconds: int = 24 * 3600) -> None:
    """清理超期暂存（上传后未发送的孤儿文件）；失败静默。"""
    import time

    staging = Path(get_settings().upload_root) / user_id
    if not staging.is_dir():
        return
    now = time.time()
    for path in staging.glob("*.json"):
        try:
            if now - path.stat().st_mtime > max_age_seconds:
                path.unlink(missing_ok=True)
        except OSError:
            continue
