"""产物工具：文档 / 海报 / 表格 / 幻灯片。

中间格式策略（决策 5/C）：文档=Markdown、海报与幻灯片=HTML、表格=xlsx+HTML 预览。
保存即：中间格式落盘 + 版本链落库（artifact_service.record_version）+
伴随最终格式文件（docx/png/xlsx/pptx，下载时直接复用）。
"""

import asyncio
import io
import json
import re
from html import escape

from app.storage.base import artifact_key
from app.tools.base import ToolRegistry, ToolSpec, ToolError
from app.tools.context import get_tool_context
from app.tools.slides import DEFAULT_THEME, resolve_theme_name

MAX_CONTENT_BYTES = 200_000  # 单产物内容上限（防滥用）


def _save(filename: str, content: str | bytes) -> str:
    """按 {user}/{task}/v{n}/{filename} 落盘（自增版本），返回存储 key。"""
    ctx = get_tool_context()
    return _save_at(ctx.next_version(), filename, content)


def _save_at(version: int, filename: str, content: str | bytes) -> str:
    """同一版本目录写文件（一次产物保存的中间格式与伴随文件同目录）。"""
    ctx = get_tool_context()
    key = artifact_key(ctx.user_id, ctx.task_id, version, filename)
    data = content.encode("utf-8") if isinstance(content, str) else content
    if len(data) > MAX_CONTENT_BYTES:
        raise ToolError(f"内容超出上限（{len(data)} > {MAX_CONTENT_BYTES} 字节）")
    ctx.storage.save(key, data)
    return key


def _record(kind: str, title: str, fmt: str, file_key: str, payload: dict | None) -> None:
    """版本链落库（无 DB 工厂的纯单测场景跳过）。"""
    from app.services.artifact_service import record_version

    ctx = get_tool_context()
    if ctx.db_factory is None:
        return
    record_version(
        ctx.db_factory,
        user_id=ctx.user_id,
        task_id=ctx.task_id,
        kind=kind,
        title=title,
        intermediate_format=fmt,
        file_key=file_key,
        preview_payload=payload,
    )


def markdown_to_docx_bytes(title: str, markdown: str) -> bytes:
    """Markdown → docx（标题层级 + 列表 + 段落；满意后下载转换复用）。"""
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=0)
    for line in markdown.splitlines():
        text = line.strip()
        if not text:
            continue
        if m := re.match(r"^#{1,4}\s+(.*)", text):
            doc.add_heading(m.group(1), level=min(len(text) - len(text.lstrip("#")), 4))
        elif re.match(r"^[-*]\s+", text):
            doc.add_paragraph(re.sub(r"^[-*]\s+", "", text), style="List Bullet")
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def _save_document(args: dict) -> str:
    title = args["title"].strip()
    content = args["content_md"].strip()
    if not content:
        raise ToolError("文档内容为空")
    from app.tools.context import get_tool_context

    v = get_tool_context().next_version()
    key = _save_at(v, "document.md", f"# {title}\n\n{content}\n")
    try:
        docx_key = _save_at(v, "document.docx", markdown_to_docx_bytes(title, content))
        extra = f"；docx={docx_key}"
    except Exception:  # 转换失败不阻塞预览（md 为主格式）
        extra = ""
    _record("document", title, "markdown", key, {"title": title})
    return json.dumps({"saved": [{"kind": "document", "key": key}], "note": f"文档《{title}》已保存{extra}"}, ensure_ascii=False)


POSTER_HTML_TEMPLATE = """<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>{title}</title>
<style>
* {{ margin: 0; padding: 0; box-sizing: border-box; }}
body {{ width: 750px; height: 1060px; font-family: "PingFang SC", "Microsoft YaHei", sans-serif;
       background: {bg}; color: {fg}; display: flex; flex-direction: column;
       align-items: center; justify-content: center; text-align: center; padding: 60px; }}
h1 {{ font-size: 64px; margin-bottom: 28px; letter-spacing: 4px; }}
.slogan {{ font-size: 28px; opacity: .9; margin-bottom: 18px; }}
.footer {{ font-size: 18px; opacity: .7; margin-top: 48px; }}
</style></head>
<body>
<h1>{title}</h1>
<p class="slogan">{slogan}</p>
<p class="footer">{footer}</p>
</body></html>
"""


async def _save_poster(args: dict) -> str:
    title = args["title"].strip()
    html = POSTER_HTML_TEMPLATE.format(
        title=escape(title),
        slogan=escape(args.get("slogan", "")),
        footer=escape(args.get("footer", "")),
        bg=args.get("bg_color", "#1f2d3d"),
        fg=args.get("fg_color", "#ffffff"),
    )
    key = _save("poster.html", html)
    png_note = ""
    try:
        png_key = await asyncio.to_thread(_render_png, key, "poster.png")
        png_note = f"；png={png_key}"
    except Exception as exc:  # 渲染器缺失/失败不阻塞，HTML 预览兜底
        png_note = f"；png渲染跳过（{type(exc).__name__}）"
    _record("poster", title, "html", key, {"title": title})
    return json.dumps({"saved": [{"kind": "poster", "key": key}], "note": f"海报《{title}》已保存{png_note}"}, ensure_ascii=False)


def _render_png(html_key: str, filename: str, width: int = 750, height: int = 1060) -> str:
    """playwright 截图（同步 API 放线程执行；无浏览器时抛错由调用方降级）。"""
    from playwright.sync_api import sync_playwright

    ctx = get_tool_context()
    html = ctx.storage.read_text(html_key)
    version_dir = html_key.rsplit("/", 1)[0]
    png_key = f"{version_dir}/{filename}"
    with sync_playwright() as p:
        browser = p.chromium.launch()
        try:
            page = browser.new_page(viewport={"width": width, "height": height})
            page.set_content(html)
            png = page.screenshot(full_page=False)
        finally:
            browser.close()
    ctx.storage.save(png_key, png)
    return png_key


async def _save_table(args: dict) -> str:
    import io

    from openpyxl import Workbook

    title = args["title"].strip()
    headers = args["headers"]
    rows = args["rows"]
    if not headers or any(len(r) != len(headers) for r in rows):
        raise ToolError("表格每行列数必须与表头一致")

    wb = Workbook()
    ws = wb.active
    # openpyxl 工作表名禁用 /\*?:[] 等字符，统一清洗
    ws.title = re.sub(r"[\\/*?:\[\]]", "-", title)[:31] or "Sheet"
    ws.append(headers)
    for r in rows:
        ws.append([str(c) for c in r])
    buf = io.BytesIO()
    wb.save(buf)
    from app.tools.context import get_tool_context

    v = get_tool_context().next_version()
    xlsx_key = _save_at(v, "table.xlsx", buf.getvalue())

    # HTML 预览数据（G8 预览接口直接可用）
    thead = "".join(f"<th>{escape(str(h))}</th>" for h in headers)
    tbody = "".join(
        "<tr>" + "".join(f"<td>{escape(str(c))}</td>" for c in r) + "</tr>" for r in rows
    )
    html = (
        "<!DOCTYPE html><html><head><meta charset='utf-8'><style>"
        "table{border-collapse:collapse;font:14px sans-serif}th,td{border:1px solid #ccc;"
        "padding:6px 14px}th{background:#f5f7fa}</style></head>"
        f"<body><h3>{escape(title)}</h3><table><thead><tr>{thead}</tr></thead>"
        f"<tbody>{tbody}</tbody></table></body></html>"
    )
    html_key = _save_at(v, "table_preview.html", html)
    _record(
        "table",
        title,
        "html_table",
        html_key,
        {"title": title, "headers": [str(h) for h in headers], "rows": [[str(c) for c in r] for r in rows]},
    )
    return json.dumps(
        {"saved": [{"kind": "table", "key": xlsx_key}, {"kind": "table_preview", "key": html_key}],
         "note": f"表格《{title}》已保存（{len(rows)} 行）"},
        ensure_ascii=False,
    )


SLIDE_THEMES = "midnight|teal|forest|coral|charcoal|berry"


async def _save_slides(args: dict) -> str:
    title = args["title"].strip()
    slides = args["slides"]
    if not slides:
        raise ToolError("至少需要一页幻灯片")
    theme = args.get("theme") or DEFAULT_THEME
    from app.tools.context import get_tool_context
    from app.tools.slides import build_slides_html, build_slides_pptx

    v = get_tool_context().next_version()
    html_key = _save_at(v, "slides.html", build_slides_html(title, slides, theme))

    pptx_key = ""
    try:
        pptx_key = _save_at(v, "slides.pptx", build_slides_pptx(title, slides, theme))
    except Exception:
        pass  # pptx 转换为尽力而为；HTML 预览为主格式（保真度降级路径）
    note = f"幻灯片《{title}》（主题 {resolve_theme_name(theme)}，{len(slides)} 页）已保存"
    if pptx_key:
        note += f"；pptx={pptx_key}"
    _record(
        "slides",
        title,
        "html_slides",
        html_key,
        {"title": title, "slides": slides},
    )
    return json.dumps({"saved": [{"kind": "slides", "key": html_key}], "note": note}, ensure_ascii=False)


def register_artifact_tools(registry: ToolRegistry) -> None:
    registry.register(
        ToolSpec(
            name="save_document",
            description="保存/更新文档产物（Markdown）。content_md 为正文（不含大标题）",
            parameters={
                "type": "object",
                "properties": {
                    "title": {"type": "string"},
                    "content_md": {"type": "string"},
                },
                "required": ["title", "content_md"],
            },
            func=_save_document,
        )
    )
    registry.register(
        ToolSpec(
            name="save_poster",
            description="保存海报产物（HTML 模板渲染，自动转 png）",
            parameters={
                "type": "object",
                "properties": {
                    "title": {"type": "string"},
                    "slogan": {"type": "string"},
                    "footer": {"type": "string"},
                    "bg_color": {"type": "string"},
                    "fg_color": {"type": "string"},
                },
                "required": ["title"],
            },
            func=_save_poster,
        )
    )
    registry.register(
        ToolSpec(
            name="save_table",
            description="保存表格产物（xlsx + HTML 预览）。rows 每行列数与 headers 一致",
            parameters={
                "type": "object",
                "properties": {
                    "title": {"type": "string"},
                    "headers": {"type": "array", "items": {"type": "string"}},
                    "rows": {
                        "type": "array",
                        "items": {"type": "array", "items": {}},
                    },
                },
                "required": ["title", "headers", "rows"],
            },
            func=_save_table,
        )
    )
    registry.register(
        ToolSpec(
            name="save_slides",
            description=(
                "保存幻灯片产物（HTML 预览 + pptx）。"
                f"theme 按内容气质选：{SLIDE_THEMES}；"
                "slides 每页 {title, bullets[], layout?, subtitle?}，"
                "layout 取 cover|agenda|section|content|end（默认 content）；"
                "封面/结尾缺省时自动补齐。数据句式规则：要点写成「数值 简短说明」"
                "（如「37% 复购率提升」）会被自动放大成数字卡片；整页都是数据要点时"
                "渲染为 KPI 大数字网格。封面 bullets 填 2-3 条数字钩子"
                "（如「120家 三年新开门店」）会显示为大数字标签；"
                "内容页 subtitle 用作 6-10 字眉题（章节归属，如「市场机会」）；"
                "总页数 ≥6 时第 2 页用 layout=agenda 放目录（bullets=各章标题）"
            ),
            parameters={
                "type": "object",
                "properties": {
                    "title": {"type": "string"},
                    "theme": {"type": "string", "enum": SLIDE_THEMES.split("|")},
                    "slides": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "title": {"type": "string"},
                                "subtitle": {"type": "string"},
                                "layout": {
                                    "type": "string",
                                    "enum": ["cover", "agenda", "section", "content", "end"],
                                },
                                "bullets": {"type": "array", "items": {"type": "string"}},
                            },
                            "required": ["title"],
                        },
                    },
                },
                "required": ["title", "slides"],
            },
            func=_save_slides,
        )
    )
